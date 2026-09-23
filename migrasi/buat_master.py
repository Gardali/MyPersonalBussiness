"""Membuat file DATABASE MASTER Posetive Photobooth (.xlsx) dari data lama.

Sumber (dibaca 23/09/2026):
- REKAP POSETIVE PHOTOBOOTH (ORDER, KAS, INVENTARIS, LOG ALAT)
- List Event / Pipeline (tab pipeline, open booth, parameter skema)
- Laporan event Lapangan UNEJ 22/09/2026 (Standar Biaya, RAB, Laporan Produksi, Laporan Keuangan)

Hanya data NYATA yang dipindah. Baris bertanda [Data Contoh] / DATA UJI SKEMA tidak ikut.
Aturan tab data: 1 baris = 1 catatan, baris 1 = judul kolom, tanpa sel gabung, tanpa rumus.
"""
import sys
from datetime import date, time
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.utils import get_column_letter

D = date
RP = '"Rp"#,##0'

wb = Workbook()
HEAD_FILL = PatternFill("solid", fgColor="1F3A5F")
HEAD_FONT = Font(bold=True, color="FFFFFF")
FLAG_FILL = PatternFill("solid", fgColor="FFF2CC")


def sheet(title, cols, rows, money=(), dates=(), times=(), pct=(), first=False):
    ws = wb.active if first else wb.create_sheet()
    ws.title = title
    ws.append([c for c, _ in cols])
    for r in rows:
        ws.append([r.get(c) for c, _ in cols])
    for i, (c, w) in enumerate(cols, 1):
        L = get_column_letter(i)
        ws.column_dimensions[L].width = w
        cell = ws.cell(row=1, column=i)
        cell.fill, cell.font = HEAD_FILL, HEAD_FONT
        cell.alignment = Alignment(vertical="center", wrap_text=True)
        fmt = RP if c in money else "yyyy-mm-dd" if c in dates else "hh:mm" if c in times else "0.0%" if c in pct else None
        if fmt:
            for row in range(2, 101):
                ws.cell(row=row, column=i).number_format = fmt
    ws.freeze_panes = "B2"
    # tandai baris yang perlu dicek (kolom catatan berisi "CEK:")
    if "catatan" in [c for c, _ in cols]:
        ci = [c for c, _ in cols].index("catatan") + 1
        for row in range(2, ws.max_row + 1):
            v = ws.cell(row=row, column=ci).value
            if v and "CEK:" in str(v):
                for col in range(1, len(cols) + 1):
                    ws.cell(row=row, column=col).fill = FLAG_FILL
    return ws


def dropdown(ws, col_name, list_range, cols):
    idx = [c for c, _ in cols].index(col_name) + 1
    L = get_column_letter(idx)
    dv = DataValidation(type="list", formula1=f"={list_range}", allow_blank=True)
    ws.add_data_validation(dv)
    dv.add(f"{L}2:{L}300")


# ---------------------------------------------------------------- PETUNJUK
ws = wb.active
ws.title = "PETUNJUK"
petunjuk = [
    ["DATABASE MASTER — POSETIVE PHOTOBOOTH"],
    [""],
    ["Satu-satunya tempat data. Sejak tanggal pindah, jangan mencatat di file lama."],
    [""],
    ["ATURAN"],
    ["1. Satu tab = satu jenis data. Satu baris = satu catatan."],
    ["2. Baris 1 adalah judul kolom. Jangan diubah, jangan disisipi baris kosong, jangan gabung sel."],
    ["3. Setiap catatan punya kode: PSV-001 (pipeline), EV-001 (event), KS-0001 (kas), ALT-001 (alat)."],
    ["4. Tab data hanya berisi isian, tanpa rumus. Total, untung/rugi, dan titik impas dihitung aplikasi."],
    ["5. Tanggal selalu format TTTT-BB-HH (contoh 2026-09-22)."],
    ["6. Baris berlatar kuning = perlu dicek (lihat kolom catatan yang diawali 'CEK:')."],
    [""],
    ["TAB"],
    ["PENGATURAN — harga jual, biaya standar, parameter skema kerja sama. Ubah angka di sini saja."],
    ["PIPELINE — calon klien dari pitching sampai deal/kalah."],
    ["EVENT — semua event (kontrak klien, open booth mandiri, kerja sama panitia)."],
    ["LAPORAN_EVENT — hasil produksi tiap event (pengganti satu file per event)."],
    ["KAS — semua uang masuk/keluar, dikaitkan ke kode event bila ada."],
    ["INVENTARIS — alat & bahan habis pakai beserta pemiliknya."],
    ["PEMAKAIAN_ALAT — alat yang dibawa & kembali per event."],
    ["LOG_KEPUTUSAN — keputusan penting, asumsi, kill criteria."],
    ["PILIHAN — isi dropdown. Tambah pilihan baru di sini."],
    ["CATATAN_MIGRASI — hal yang perlu dicek dari pemindahan data lama."],
]
for r in petunjuk:
    ws.append(r)
ws["A1"].font = Font(bold=True, size=14)
for c in ("A5", "A13"):
    ws[c].font = Font(bold=True)
ws.column_dimensions["A"].width = 110

# ---------------------------------------------------------------- PILIHAN
pilihan = {
    "Sumber": ["Tawaran Masuk", "Pitching", "Open Booth", "Referensi"],
    "Kanal": ["WhatsApp", "Instagram", "WO/EO", "Vendor Lain", "Relasi Pribadi", "Kampus/Organisasi", "Lainnya"],
    "Jenis Acara": ["Pernikahan", "Lamaran", "Ulang Tahun", "Wisuda", "Acara Kampus", "Korporat/Gathering",
                    "Aktivasi Brand", "Bazar/Festival", "Lainnya"],
    "Model Pendapatan": ["Kontrak Klien", "Open Booth Mandiri", "Kerja Sama Panitia"],
    "Skema": ["A. Berjenjang", "B. Persentase Omzet", "C. Di Atas Impas", "D. Fee Tetap",
              "E. Voucher Prabayar", "F. Bundling Peserta"],
    "Status Pipeline": ["Prospek", "Sudah Dihubungi", "Penawaran Dikirim", "Negosiasi", "Deal", "Kalah", "Batal"],
    "Alasan Kalah": ["Harga Terlalu Mahal", "Tanggal Bentrok", "Pilih Vendor Lain", "Acara Batal",
                     "Tidak Ada Respons", "Lokasi Tidak Terjangkau", "Lainnya"],
    "Status Event": ["Terkonfirmasi", "Selesai", "Batal"],
    "Jenis Kas": ["Masuk", "Keluar"],
    "Kategori Kas": ["Penjualan Event", "DP / Pelunasan Kontrak", "Setoran Owner", "Pengadaan Barang",
                     "Media Cetak", "Fee Crew", "Transport", "Konsumsi", "Fee Panitia / Sewa Lapak",
                     "Langganan Software", "Lainnya"],
    "Metode": ["Tunai", "QRIS", "Transfer", "Dibayar Owner Langsung"],
    "Kategori Alat": ["Aset", "Barang Habis Pakai"],
    "Pemilik": ["Posetive", "Mas Bagus (Sinektive)", "Pribadi Garda"],
    "Kondisi": ["Baik", "Perlu Perbaikan", "Rusak", "Hilang"],
    "Ya/Tidak": ["Ya", "Tidak"],
    "Indoor/Outdoor": ["Indoor", "Outdoor"],
}
wp = wb.create_sheet("PILIHAN")
ranges = {}
for i, (k, vals) in enumerate(pilihan.items(), 1):
    L = get_column_letter(i)
    wp.cell(row=1, column=i, value=k).font = HEAD_FONT
    wp.cell(row=1, column=i).fill = HEAD_FILL
    for j, v in enumerate(vals, 2):
        wp.cell(row=j, column=i, value=v)
    wp.column_dimensions[L].width = 24
    ranges[k] = f"PILIHAN!${L}$2:${L}$40"

# ---------------------------------------------------------------- PENGATURAN
set_cols = [("kunci", 26), ("nilai", 14), ("satuan", 12), ("keterangan", 70)]
S = lambda k, v, s, ket="": {"kunci": k, "nilai": v, "satuan": s, "keterangan": ket}
pengaturan = [
    S("HARGA_KERTAS", 70000, "Rp/pack", "Kertas foto Imapro 4R"),
    S("ISI_KERTAS", 100, "lembar", "Isi per pack"),
    S("HARGA_TINTA", 1300000, "Rp/set", "Tinta Epson 057, 6 warna"),
    S("KAPASITAS_TINTA", 2100, "lembar", "Klaim Epson — perbarui dari data nyata"),
    S("CADANGAN_GAGAL", 0.05, "persen", "Cadangan cetak gagal"),
    S("BIAYA_BUKU", 1930, "Rp/buku", "Biaya cetak Graduation Book, sesuai nota"),
    S("LEMBAR_PER_STRIP", 1, "lembar", "1 lembar = 2 strip, dijual per lembar"),
    S("FEE_CREW", 60000, "Rp/orang", "Fee crew per orang per event"),
    S("DURASI_TERCAKUP", 6, "jam", "Durasi kerja yang tercakup fee crew"),
    S("TARIF_LEMBUR", 10000, "Rp/jam", "Lembur per orang per jam"),
    S("TRANSPORT_STD", 50000, "Rp/event", "Transport PP dalam kota"),
    S("KONSUMSI_ORANG", 25000, "Rp/orang", ""),
    S("KAS_KEMBALIAN_STD", 20000, "Rp", "Modal kembalian bila menerima tunai"),
    S("HARGA_4R", 20000, "Rp/lembar", "Harga jual foto 4R"),
    S("HARGA_BUKU", 35000, "Rp/buku", "Harga jual Graduation Book (termasuk 1 foto 4R)"),
    S("HARGA_STRIP", 10000, "Rp/lembar", "Cetak strip tambahan"),
    S("POTONGAN_GRUP", 5000, "Rp/grup", "Dari total, bukan per orang"),
    S("JEPRETO_BULAN", 600000, "Rp/bulan", "Langganan software Jepreto"),
    S("PENYUSUTAN_BULAN", 325584, "Rp/bulan", "SEMENTARA: aset Posetive ≥ Rp250.000; alat Sinektive tidak termasuk"),
    S("EVENT_BULAN", 4, "event", "Target event per bulan (dasar alokasi Jepreto & penyusutan)"),
    S("TOLERANSI", 0.10, "persen", "Batas toleransi selisih biaya RAB vs realisasi"),
    S("SKEMA_A_JENJANG", "0:0;15:3000;45:4500;70:5500;100:7000", "sesi:Rp", "A. Berjenjang — sesi minimal : tarif ke panitia per sesi"),
    S("SKEMA_B_PERSEN", 0.25, "persen", "B. Persentase omzet untuk panitia"),
    S("SKEMA_C_BATAS", 30, "sesi", "C. Batas sesi untuk Posetive sebelum bagi hasil"),
    S("SKEMA_C_BAGIAN", 8000, "Rp/sesi", "C. Bagian panitia per sesi di atas batas"),
    S("SKEMA_D_FEE", 300000, "Rp", "D. Fee tetap default"),
    S("SKEMA_E_HARGA", 15000, "Rp/voucher", "E. Harga voucher ke panitia"),
    S("SKEMA_E_MIN", 40, "voucher", "E. Minimal voucher"),
    S("SKEMA_F_HARGA", 12000, "Rp/peserta", "F. Harga bundling per peserta"),
    S("SKEMA_F_MIN", 60, "peserta", "F. Minimal peserta"),
]
sheet("PENGATURAN", set_cols, pengaturan)

# ---------------------------------------------------------------- PIPELINE
pipe_cols = [("id", 10), ("tanggal_masuk", 13), ("sumber", 14), ("kanal", 16), ("referensi", 22),
             ("nama_klien", 26), ("kontak", 16), ("jenis_acara", 18), ("tanggal_event", 13), ("kota", 12),
             ("luar_kota", 9), ("indoor_outdoor", 11), ("estimasi_tamu", 10), ("paket", 10),
             ("model_pendapatan", 20), ("skema", 20), ("parameter_skema", 12), ("estimasi_sesi", 10),
             ("harga_penawaran", 14), ("harga_deal", 14), ("status", 14), ("alasan_kalah", 20),
             ("follow_up", 13), ("id_event", 10), ("catatan", 60)]
pipeline = [
    dict(id="PSV-001", tanggal_masuk=D(2026, 9, 22), sumber="Pitching", kanal="WhatsApp", referensi="Inaugurasi FIB",
         nama_klien="Rasya", kontak="6281232458383", jenis_acara="Acara Kampus", tanggal_event=D(2026, 10, 10),
         kota="Jember", luar_kota="Tidak", indoor_outdoor="Outdoor", estimasi_tamu=500, paket="Paket B",
         model_pendapatan="Kontrak Klien", status="Prospek", follow_up=D(2026, 10, 1),
         catatan="Follow up ke Rasya untuk memastikan ketersediaan photobooth"),
]
wsp = sheet("PIPELINE", pipe_cols, pipeline, money=("harga_penawaran", "harga_deal"),
            dates=("tanggal_masuk", "tanggal_event", "follow_up"))
for col, lst in [("sumber", "Sumber"), ("kanal", "Kanal"), ("jenis_acara", "Jenis Acara"),
                 ("model_pendapatan", "Model Pendapatan"), ("skema", "Skema"), ("status", "Status Pipeline"),
                 ("alasan_kalah", "Alasan Kalah"), ("luar_kota", "Ya/Tidak"), ("indoor_outdoor", "Indoor/Outdoor")]:
    dropdown(wsp, col, ranges[lst], pipe_cols)

# ---------------------------------------------------------------- EVENT
ev_cols = [("id_event", 10), ("id_pipeline", 10), ("nama_event", 24), ("tanggal", 12), ("jam_buka", 8),
           ("jam_tutup", 8), ("lokasi", 22), ("kota", 12), ("jenis_acara", 16), ("model_pendapatan", 20),
           ("skema", 18), ("parameter_skema", 12), ("nilai_kontrak", 14), ("jumlah_crew", 8), ("status", 13),
           ("penanggung_jawab", 18), ("tim", 24), ("folder_drive", 30), ("catatan", 60)]
events = [
    dict(id_event="EV-001", nama_event="Wisuda Unej", tanggal=D(2026, 9, 19), lokasi="Universitas Jember",
         kota="Jember", jenis_acara="Wisuda", status="Terkonfirmasi", tim="Garda, Daffa, Syafa",
         folder_drive="2026-09-19_WisudaUnej_Wisuda",
         catatan="CEK: tanggal sudah lewat tapi status masih Terkonfirmasi & belum ada laporan. Model pendapatan belum diisi."),
    dict(id_event="EV-002", nama_event="Lapangan UNEJ", tanggal=D(2026, 9, 22), jam_buka=time(14, 30),
         jam_tutup=time(17, 50), lokasi="Lapangan UNEJ", kota="Jember", jenis_acara="Bazar/Festival",
         model_pendapatan="Open Booth Mandiri", jumlah_crew=1, status="Selesai", penanggung_jawab="Garda Ali Rayhaan",
         tim="Garda Ali Rayhaan, Daffa Falih Naufal",
         catatan="CEK: tidak tercatat di ORDER lama. RAB tertulis 1 crew, laporan produksi mencatat 2 orang bertugas. Jenis acara diisi Bazar/Festival — ganti bila kurang tepat."),
    dict(id_event="EV-003", nama_event="Wisuda Unej", tanggal=D(2026, 9, 26), lokasi="Universitas Jember",
         kota="Jember", jenis_acara="Wisuda", status="Terkonfirmasi", tim="Garda, Daffa, Syafa",
         folder_drive="2026-09-26_WisudaUnej_Wisuda", catatan="CEK: model pendapatan belum diisi."),
]
wse = sheet("EVENT", ev_cols, events, money=("nilai_kontrak",), dates=("tanggal",), times=("jam_buka", "jam_tutup"))
for col, lst in [("jenis_acara", "Jenis Acara"), ("model_pendapatan", "Model Pendapatan"), ("skema", "Skema"),
                 ("status", "Status Event")]:
    dropdown(wse, col, ranges[lst], ev_cols)

# ---------------------------------------------------------------- LAPORAN_EVENT
lap_cols = [("id_event", 10), ("jam_buka_aktual", 9), ("jam_tutup_aktual", 9), ("jumlah_transaksi", 10),
            ("foto_4r_terjual", 10), ("grad_book_terjual", 10), ("strip_terjual", 10), ("grup_potongan", 10),
            ("counter_awal", 10), ("counter_akhir", 10), ("lembar_uji_bonus", 10), ("lembar_gagal", 10),
            ("softfile_terkirim", 10), ("tunai_dihitung", 13), ("modal_kembalian", 13), ("qris_transfer", 13),
            ("stok_kertas_awal", 10), ("stok_kertas_akhir_fisik", 10), ("stok_buku_awal", 10), ("buku_rusak", 10),
            ("tinta_terendah_persen", 10), ("jam_ramai", 14), ("antrean_terpanjang", 10), ("downtime_menit", 10),
            ("kendala", 30), ("keluhan_pengunjung", 30), ("rekomendasi", 50), ("catatan", 50)]
laporan = [
    dict(id_event="EV-002", jam_buka_aktual=time(14, 10), jam_tutup_aktual=time(17, 50), jumlah_transaksi=3,
         foto_4r_terjual=3, grad_book_terjual=0, strip_terjual=0, grup_potongan=0, counter_awal=0, counter_akhir=5,
         lembar_uji_bonus=2, lembar_gagal=0, softfile_terkirim=3, tunai_dihitung=40000, modal_kembalian=20000,
         qris_transfer=40000, stok_kertas_awal=200, stok_kertas_akhir_fisik=197, stok_buku_awal=25, buku_rusak=0,
         tinta_terendah_persen=1.0, downtime_menit=0,
         rekomendasi="Crew pas 2-3 orang. Jam buka terbaik 16:00 kalau di lapangan UNEJ. Stok dibawa 100 lembar.",
         catatan="CEK: stok kertas awal 200 lembar, tapi INVENTARIS mencatat 5 pack (500 lembar). Hitungan akhir 195 vs fisik 197 (selisih 2)."),
]
sheet("LAPORAN_EVENT", lap_cols, laporan, money=("tunai_dihitung", "modal_kembalian", "qris_transfer"),
      times=("jam_buka_aktual", "jam_tutup_aktual"), pct=("tinta_terendah_persen",))

# ---------------------------------------------------------------- KAS
kas_cols = [("id", 9), ("tanggal", 12), ("jenis", 8), ("kategori", 20), ("keterangan", 42), ("id_event", 9),
            ("nominal", 14), ("metode", 20), ("no_nota", 16), ("catatan", 60)]
K = []


def kas(tgl, jenis, kat, ket, nominal, metode=None, nota=None, ev=None, cat=None):
    K.append(dict(id=f"KS-{len(K) + 1:04d}", tanggal=tgl, jenis=jenis, kategori=kat, keterangan=ket, id_event=ev,
                  nominal=nominal, metode=metode, no_nota=nota, catatan=cat))


t14, t17, t22 = D(2026, 9, 14), D(2026, 9, 17), D(2026, 9, 22)
for ket, n, nota in [("Background", 186670, "01.2026.09.01"), ("Meja", 416949, "01.2026.09.02"),
                     ("Power Station", 1480296, "01.2026.09.03"), ("Stempel Tanggal", 21776, "01.2026.09.04")]:
    kas(t14, "Masuk", "Setoran Owner", f"Dana owner untuk beli {ket}", n, "Dibayar Owner Langsung")
    kas(t14, "Keluar", "Pengadaan Barang", f"Beli {ket} (langsung dibeli oleh owner)", n, "Dibayar Owner Langsung", nota)
kas(t17, "Masuk", "Setoran Owner", "Dana owner untuk beli prop", 250000)
kas(t17, "Keluar", "Pengadaan Barang", "Beli Prop (Blink-Blink)", 140900, nota="01.2026.09.05.pdf")
kas(t17, "Keluar", "Pengadaan Barang", "Beli Prop (Tuku-tuku)", 79000, nota="01.2026.09.06.pdf")
kas(t17, "Masuk", "Setoran Owner", "Dana owner untuk beli Roll Banner", 250000)
kas(t17, "Keluar", "Pengadaan Barang", "Beli Roll Banner", 250000, nota="01.2026.09.07")
cek = "CEK: dari laporan event 22/09, belum tercatat di KAS lama. Pastikan nominal & dari mana uangnya."
kas(t22, "Masuk", "Penjualan Event", "Penjualan tunai Lapangan UNEJ (3 lembar 4R, sebagian)", 20000, "Tunai", ev="EV-002", cat=cek)
kas(t22, "Masuk", "Penjualan Event", "Penjualan QRIS Lapangan UNEJ", 40000, "QRIS", ev="EV-002", cat=cek)
kas(t22, "Keluar", "Fee Crew", "Fee crew Lapangan UNEJ", 50000, ev="EV-002", cat=cek)
kas(t22, "Keluar", "Transport", "Transport Lapangan UNEJ", 46000, ev="EV-002", cat=cek)
kas(t22, "Keluar", "Konsumsi", "Konsumsi Lapangan UNEJ (Gojek)", 40000, ev="EV-002", cat=cek)
wsk = sheet("KAS", kas_cols, K, money=("nominal",), dates=("tanggal",))
for col, lst in [("jenis", "Jenis Kas"), ("kategori", "Kategori Kas"), ("metode", "Metode")]:
    dropdown(wsk, col, ranges[lst], kas_cols)

# ---------------------------------------------------------------- INVENTARIS
inv_cols = [("id_alat", 9), ("nama", 30), ("kategori", 18), ("pemilik", 22), ("jumlah", 8), ("satuan", 8),
            ("kondisi", 14), ("lokasi", 14), ("tanggal_masuk", 13), ("nilai_beli", 14), ("catatan", 50)]
I = []


def alat(nama, kat, pemilik, jml, satuan, kondisi, lokasi, tgl, nilai, cat=None):
    I.append(dict(id_alat=f"ALT-{len(I) + 1:03d}", nama=nama, kategori=kat, pemilik=pemilik, jumlah=jml,
                  satuan=satuan, kondisi=kondisi, lokasi=lokasi, tanggal_masuk=tgl, nilai_beli=nilai, catatan=cat))


MB, PS, KG = "Mas Bagus (Sinektive)", "Posetive", "Kos Garda"
fix = "Tanggal masuk diperbaiki (di file lama melompat ke 2027-2035 karena autofill)."
alat("Kamera Nikon", "Aset", MB, 1, "unit", "Baik", KG, t14, None, fix)
alat("Lighting TT600", "Aset", MB, 1, "unit", "Baik", KG, t14, None, fix)
alat("Printer Epson L8050", "Aset", PS, 1, "unit", "Baik", KG, t14, 4556000, fix)
alat("Meja Cutting", "Aset", PS, 1, "unit", "Baik", KG, t14, 145652, fix)
alat("Layar Goojodoq", "Aset", PS, 1, "unit", "Baik", KG, t14, 1427528, fix)
alat("Kertas Imapro 4R @100", "Barang Habis Pakai", PS, 5, "pack", "Baik", KG, t14, 70000,
     "CEK: 5 pack = 500 lembar, laporan 22/09 mencatat stok awal 200 lembar. Nilai = harga per pack.")
alat("Kertas Imapro A4 @20", "Barang Habis Pakai", PS, 10, "pack", "Baik", KG, t14, 50000, "Nilai = harga per pack. " + fix)
alat("Baterai Nikon", "Aset", MB, 1, "unit", "Baik", KG, t14, None, fix)
alat("Charger Baterai Nikon", "Aset", MB, 1, "unit", "Baik", KG, t14, None, fix)
alat("Strap Kamera Nikon", "Aset", MB, 1, "unit", "Baik", KG, t14, None, fix)
alat("Background Biru Abstrak", "Aset", PS, 1, "unit", "Baik", KG, t14, 186670, "Tanggal dari KAS")
alat("Meja Lipat Hitam", "Aset", PS, 1, "unit", "Baik", KG, t14, 416949, "Tanggal dari KAS")
alat("Power Station 220V 400W-1200W", "Aset", PS, 1, "unit", "Baik", KG, t14, 1480296, "Tanggal dari KAS")
alat("Joyko Stempel Tanggal", "Aset", PS, 1, "unit", "Baik", KG, t14, 21776, "Tanggal dari KAS")
alat("Flash Godox MS300V", "Aset", PS, 1, "unit", None, None, None, 1456800, "CEK: kondisi, lokasi, tanggal masuk kosong")
alat("HDMI to Mini HDMI", "Aset", PS, 1, "unit", "Baik", None, None, 78400)
alat("Stand Backdrop", "Aset", PS, 1, "unit", None, None, None, 195011, "CEK: kondisi kosong")
alat("Ball Head", "Aset", PS, 2, "unit", "Baik", None, None, 19816, "@Rp9.908")
alat("Adapter Type C Dongle", "Aset", PS, 1, "unit", "Baik", None, None, 82207)
alat("HDMI Video Capture", "Aset", PS, 1, "unit", "Baik", None, None, 83751)
alat("Converter HDMI", "Aset", PS, 1, "unit", "Baik", None, None, 26129)
alat("OTG USB to Type C", "Aset", PS, 1, "unit", "Baik", None, None, 16379)
alat("Extension USB", "Aset", PS, 1, "unit", "Baik", None, None, 55407)
alat("Bracket Monitor", "Aset", PS, 1, "unit", "Baik", None, None, 188600)
alat("Cable Time", "Aset", PS, 1, "unit", "Baik", None, None, 23875)
alat("Cable Data Nikon", "Aset", PS, 1, "unit", "Baik", None, None, 37100)
alat("Payung Flash", "Aset", PS, 1, "unit", None, None, None, 39990, "CEK: kondisi kosong")
alat("Custom Box", "Aset", PS, 1, "unit", None, None, None, 2000000, "CEK: kondisi kosong")
alat("Baterai Dummy Kamera Nikon D7100", "Aset", PS, 1, "unit", "Baik", None, None, 255600)
alat("Roll Banner", "Aset", PS, 1, "unit", "Baik", None, t17, 250000, "Tanggal dari KAS")
alat("Prop Kacamata", "Aset", PS, 9, "pcs", "Baik", KG, t17, 140900, "Nilai Rp140.900 adalah gabungan kacamata + bandana (nota Blink-Blink)")
alat("Prop Bandana", "Aset", PS, 2, "pcs", "Baik", KG, t17, None, "Termasuk dalam nota Blink-Blink bersama kacamata")
alat("Prop Bando", "Aset", PS, 5, "pcs", "Baik", KG, t17, 79000)
alat("Mini PC / Laptop", "Aset", PS, 1, "unit", "Baik", None, None, None,
     "CEK: dibawa ke event 22/09 tapi tidak ada di inventaris lama. Isi nilai & pemilik.")
wsi = sheet("INVENTARIS", inv_cols, I, money=("nilai_beli",), dates=("tanggal_masuk",))
for col, lst in [("kategori", "Kategori Alat"), ("pemilik", "Pemilik"), ("kondisi", "Kondisi")]:
    dropdown(wsi, col, ranges[lst], inv_cols)

# ---------------------------------------------------------------- PEMAKAIAN_ALAT
by_name = {a["nama"]: a["id_alat"] for a in I}
pa_cols = [("id_event", 10), ("id_alat", 10), ("nama_alat", 30), ("dibawa", 8), ("kembali", 8),
           ("kondisi_akhir", 14), ("catatan", 40)]
bawa = [("Printer Epson L8050", True), ("Kamera Nikon", True), ("Flash Godox MS300V", False),
        ("Lighting TT600", True), ("Layar Goojodoq", True), ("Power Station 220V 400W-1200W", True),
        ("Custom Box", True), ("Mini PC / Laptop", True), ("Stand Backdrop", True),
        ("Background Biru Abstrak", False), ("Meja Lipat Hitam", True), ("Set prop", True),
        ("Set kabel & adapter", True), ("Payung Flash", True), ("Baterai Dummy Kamera Nikon D7100", True)]
pa = [dict(id_event="EV-002", id_alat=by_name.get(n), nama_alat=n, dibawa="Ya" if b else "Tidak",
           kembali="Ya" if b else None, kondisi_akhir="Baik",
           catatan=None if n in by_name else "Set (beberapa alat) — tidak punya kode tunggal") for n, b in bawa]
wsa = sheet("PEMAKAIAN_ALAT", pa_cols, pa)
for col, lst in [("dibawa", "Ya/Tidak"), ("kembali", "Ya/Tidak"), ("kondisi_akhir", "Kondisi")]:
    dropdown(wsa, col, ranges[lst], pa_cols)

# ---------------------------------------------------------------- LOG_KEPUTUSAN
sheet("LOG_KEPUTUSAN", [("tanggal", 12), ("keputusan", 40), ("asumsi", 40), ("kill_criteria", 40),
                        ("tanggal_tinjau", 13), ("hasil_tinjauan", 40)], [], dates=("tanggal", "tanggal_tinjau"))

# ---------------------------------------------------------------- CATATAN_MIGRASI
cm_cols = [("no", 5), ("tab", 16), ("hal", 70), ("tindakan", 60), ("status", 12)]
cm = [
    ("EVENT", "Wisuda Unej 19/09 sudah lewat, status masih Terkonfirmasi, tidak ada laporan.", "Apakah event jalan? Kalau ya, isi LAPORAN_EVENT & ubah status jadi Selesai."),
    ("EVENT", "Model pendapatan kedua Wisuda Unej belum diketahui.", "Pilih: Kontrak Klien / Open Booth Mandiri / Kerja Sama Panitia."),
    ("EVENT", "Event 22/09 Lapangan UNEJ tidak pernah tercatat di ORDER lama.", "Sudah ditambahkan sebagai EV-002."),
    ("KAS", "Uang penjualan & biaya event 22/09 tidak tercatat di KAS lama. Sudah dimasukkan dari laporan event (baris kuning).", "Cek: biaya Rp136.000 dibayar dari mana? Kalau dari uang pribadi, catat sebagai Setoran Owner / talangan."),
    ("KAS", "Setoran dana dari owner dulu berkategori Pengadaan Barang.", "Diubah jadi Setoran Owner supaya tidak terhitung sebagai pendapatan."),
    ("INVENTARIS", "Tanggal masuk melompat ke 2027-2035 (efek autofill).", "Diperbaiki jadi 2026-09-14."),
    ("INVENTARIS", "Stok kertas 4R: inventaris 5 pack (500 lembar) vs laporan 22/09 stok awal 200 lembar.", "Hitung fisik, lalu betulkan jumlah."),
    ("INVENTARIS", "Mini PC / Laptop dibawa ke event tapi tidak ada di inventaris.", "Ditambahkan (ALT-034). Isi pemilik & nilai."),
    ("LAPORAN_EVENT", "Laporan 22/09 menunjukkan status lembar PERIKSA (selisih 3), karena rumus 'lembar terjual' error (#NAME?).", "Hitungan benar: 5 tercetak = 3 terjual + 2 uji + 0 gagal → selisih 0, cocok."),
    ("PENGATURAN", "Biaya bahan per sesi di file pipeline Rp1.451, sedangkan dari Standar Biaya ±Rp1.385 (kertas Rp700 + tinta Rp619, +5% cadangan).", "Pakai satu sumber: tab PENGATURAN. Aplikasi menghitung dari angka dasar."),
    ("PIPELINE", "Baris [Data Contoh] dan DATA UJI SKEMA (PSV-002 s.d. PSV-009) serta 2 baris contoh open booth tidak dipindah.", "Kalau ada yang sebenarnya data nyata, beri tahu."),
]
sheet("CATATAN_MIGRASI", cm_cols, [dict(no=i, tab=t, hal=h, tindakan=a, status="Belum") for i, (t, h, a) in enumerate(cm, 1)])

# urutan tab
order = ["PETUNJUK", "CATATAN_MIGRASI", "PIPELINE", "EVENT", "LAPORAN_EVENT", "KAS", "INVENTARIS",
         "PEMAKAIAN_ALAT", "LOG_KEPUTUSAN", "PENGATURAN", "PILIHAN"]
wb._sheets = [wb[n] for n in order]

out = sys.argv[1] if len(sys.argv) > 1 else "POSETIVE_Database_Master.xlsx"
wb.save(out)
print("tersimpan:", out)
